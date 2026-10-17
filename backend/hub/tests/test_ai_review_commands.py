import io
import json
import tempfile
from pathlib import Path
from unittest import mock

from django.contrib.auth.models import User
from django.core.management import call_command
from django.test import TestCase
from openpyxl import load_workbook

from hub.ai_review import model
from hub.models import AIReview, AssignmentReviewConfig, AssignmentSubmission, UserProfile

from .test_ai_review_model import CRITERIA, answer
from .test_ai_review_models import make_assignment


class CommandTests(TestCase):
    def setUp(self):
        _, _, self.activity, self.resource = make_assignment()
        AssignmentReviewConfig.objects.create(resource=self.resource, enabled=True, criteria=CRITERIA, rubric_approved=True)
        self.subs = []
        for name, status, opt_out in (('ann', 'approved', False), ('bob', 'changes_requested', True)):
            user = User.objects.create_user(name, password='x', first_name=name.title(), email=f'{name}@example.org')
            UserProfile.objects.create(user=user, ai_training_opt_out=opt_out)
            sub = AssignmentSubmission.objects.create(user=user, lesson=self.activity, resource=self.resource,
                                                      text=f'{name.title()} wrote: bias, privacy', status=status,
                                                      feedback=f'Thanks {name.title()}')
            AIReview.objects.create(submission=sub, status='done', rubric=CRITERIA, verdict='approve', confidence=0.9,
                                    submission_text=sub.text, sent_feedback=sub.feedback,
                                    human_verdict='approve' if status == 'approved' else 'changes_requested',
                                    agreed=status == 'approved')
            self.subs.append(sub)

    def test_benchmark_writes_summary(self):
        out = Path(tempfile.mkdtemp()) / 'bench.xlsx'
        with mock.patch.object(model, '_chat', return_value=answer()):
            call_command('ai_review_benchmark', '--models', 'm1,m2', '--out', str(out), stdout=io.StringIO())
        wb = load_workbook(out)
        summary = list(wb['Summary'].values)
        self.assertEqual(summary[0][:3], ('Model', 'Submissions', 'Agreement'))
        self.assertEqual({r[0] for r in summary[1:]}, {'m1', 'm2'})
        self.assertEqual(wb['Answers'].max_row, 1 + 2 * 2)
        self.assertEqual(AIReview.objects.count(), 2)            # nothing stored

    def test_export_excludes_opted_out_and_anonymises(self):
        out = Path(tempfile.mkdtemp()) / 'train.jsonl'
        call_command('ai_review_export', '--out', str(out), stdout=io.StringIO())
        lines = [json.loads(line) for line in out.read_text(encoding='utf-8').splitlines()]
        self.assertEqual(len(lines), 1)                          # Bob opted out
        self.assertNotIn('Ann', lines[0]['submission_text'] + lines[0]['feedback'])
        self.assertEqual(lines[0]['human_verdict'], 'approve')
