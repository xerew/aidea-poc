from io import BytesIO

from django.urls import reverse
from openpyxl import load_workbook
from rest_framework import status
from rest_framework.test import APITestCase

from analytics.workbook import SHEET_NAMES, build_learning_workbook
from hub.models import Course, UserProfile

from .fixtures import NOW, CourseFixture, make_user


def rows(ws):
    """Sheet rows as dicts keyed by the header row."""
    values = list(ws.values)
    return [dict(zip(values[0], row)) for row in values[1:]]


def by_learner(ws, name):
    return [r for r in rows(ws) if r['Learner'] == name]


class WorkbookTests(APITestCase):
    @classmethod
    def setUpTestData(cls):
        cls.f = CourseFixture()

    def setUp(self):
        self.wb = build_learning_workbook([self.f.course], now=NOW)

    def test_sheet_order(self):
        self.assertEqual(self.wb.sheetnames, SHEET_NAMES)
        self.assertEqual(SHEET_NAMES, [
            'README', 'Overview', 'Learners', 'Modules', 'Activities', 'Resources',
            'Visits', 'Quiz answers', 'H5P answers', 'Events',
        ])

    def test_learners_sheet(self):
        [ada] = by_learner(self.wb['Learners'], 'Ada Byte')
        self.assertEqual(ada['Email'], 'ada@example.org')
        self.assertEqual((ada['Active s'], ada['On-screen s'], ada['Visits']), (350, 390, 2))
        self.assertEqual(ada['Status'], 'on_track')
        self.assertEqual(ada['Study participant (consented)'], 'yes')
        self.assertEqual(ada['Course ID'], self.f.course.id)
        [di] = by_learner(self.wb['Learners'], 'di')
        self.assertEqual((di['Status'], di['Study participant (consented)']), ('stuck', 'no'))

    def test_overview_is_wide_in_minutes(self):
        ws = self.wb['Overview']
        header = [c.value for c in ws[1]]
        self.assertIn('M1 Basics — active min', header)
        self.assertIn('M2 Practice — % complete', header)
        [ada] = by_learner(ws, 'Ada Byte')
        self.assertEqual(ada['M1 Basics — active min'], 5.8)   # 350 s
        self.assertEqual(ada['Total active min'], 5.8)

    def test_modules_and_activities_sheets(self):
        m1 = next(r for r in by_learner(self.wb['Modules'], 'Ada Byte') if r['Module'] == 'Basics')
        self.assertEqual((m1['Active s'], m1['Visits'], m1['Activities done'], m1['Activities total']), (350, 2, 0, 2))
        a1 = next(r for r in by_learner(self.wb['Activities'], 'Bo Bit') if r['Activity'] == 'Read and watch')
        self.assertEqual((a1['Resources done'], a1['Resources total'], a1['Done']), (2, 2, 'yes'))
        self.assertTrue(a1['Completed at'].endswith('Z'))

    def test_resources_sheet_one_row_per_learner_and_resource(self):
        ws = self.wb['Resources']
        self.assertEqual(len(rows(ws)), 6 * 5)
        text = next(r for r in by_learner(ws, 'Ada Byte') if r['Resource'] == 'Intro')
        self.assertEqual((text['Active s'], text['Scroll %'], text['Type'], text['Required']), (150, 80, 'text', 'yes'))
        pdf = next(r for r in by_learner(ws, 'Bo Bit') if r['Resource'] == 'Sheet')
        self.assertEqual((pdf['PDF opened'], pdf['PDF downloaded']), (1, 1))
        old_text = next(r for r in by_learner(ws, 'old') if r['Resource'] == 'Intro')
        # Done before tracking began: no visits, so time is blank, not 0.
        self.assertEqual((old_text['Visits'], old_text['Active s'], old_text['On-screen s']), (0, None, None))
        self.assertTrue(old_text['Completed at'])

    def test_unmeasured_time_is_blank_at_every_level(self):
        [old] = by_learner(self.wb['Learners'], 'old')
        self.assertEqual((old['Visits'], old['Active s'], old['On-screen s']), (0, None, None))
        [old_m1] = [r for r in by_learner(self.wb['Modules'], 'old') if r['Module'] == 'Basics']
        self.assertEqual((old_m1['Visits'], old_m1['Active s']), (0, None))
        [old_a1] = [r for r in by_learner(self.wb['Activities'], 'old') if r['Activity'] == 'Read and watch']
        self.assertEqual((old_a1['Visits'], old_a1['Active s']), (0, None))
        [old_o] = by_learner(self.wb['Overview'], 'old')
        self.assertEqual((old_o['M1 Basics — active min'], old_o['Total active min']), (None, None))
        [ada] = by_learner(self.wb['Learners'], 'Ada Byte')
        self.assertEqual(ada['Active s'], 350)   # measured time is still a number

    def test_text_cells_are_safe(self):
        f = self.f
        f.ada.first_name = '=HYPERLINK("http://x")'
        f.ada.save()
        f.quiz.quiz_data[0]['question'] = 'Q\x01 one'
        f.quiz.save()
        wb = build_learning_workbook([f.course], now=NOW)
        buffer = BytesIO()
        wb.save(buffer)  # control characters would raise IllegalCharacterError
        names = [r['Learner'] for r in rows(load_workbook(BytesIO(buffer.getvalue()))['Learners'])]
        self.assertIn("'=HYPERLINK(\"http://x\") Byte", names)
        questions = [r['Question'] for r in rows(load_workbook(BytesIO(buffer.getvalue()))['Quiz answers'])]
        self.assertIn('Q one', questions)

    def test_visits_sheet(self):
        visits = by_learner(self.wb['Visits'], 'Ada Byte')
        self.assertEqual(len(visits), 3)
        mobile = next(v for v in visits if v['Device'] == 'mobile')
        self.assertEqual((mobile['Language'], mobile['Local hour'], mobile['TZ offset (min)'], mobile['Completed during']),
                         ('el', 14, 180, 'yes'))
        self.assertTrue(mobile['Started at'].endswith('Z'))

    def test_quiz_answers_sheet(self):
        answers = by_learner(self.wb['Quiz answers'], 'Bo Bit')
        self.assertEqual([(a['Question #'], a['Option picked'], a['Right'], a['Seconds on question']) for a in answers],
                         [(1, 'a', 'yes', 12.5), (2, 'd', 'yes', 7.5)])

    def test_events_sheet(self):
        seek = next(e for e in by_learner(self.wb['Events'], 'Ada Byte') if e['Event'] == 'video_seek')
        self.assertEqual((seek['From s'], seek['To s']), (30, 90))

    def test_workbook_for_empty_course(self):
        empty = Course.objects.create(title='Empty', pillar=self.f.course.pillar, level='beginner', duration_hours=1)
        wb = build_learning_workbook([empty], now=NOW)
        self.assertEqual(wb.sheetnames, SHEET_NAMES)
        self.assertEqual(rows(wb['Learners']), [])


class ExportEndpointTests(APITestCase):
    @classmethod
    def setUpTestData(cls):
        cls.f = CourseFixture()
        cls.outsider = make_user('la_other', UserProfile.UserType.CONTENT_CREATOR)

    def test_course_export(self):
        self.client.force_authenticate(self.f.creator)
        res = self.client.get(reverse('analytics-course-export', kwargs={'pk': self.f.course.id}))
        self.assertEqual(res.status_code, status.HTTP_200_OK)
        self.assertIn('spreadsheetml', res['Content-Type'])
        self.assertIn('analytics-course-analytics.xlsx', res['Content-Disposition'])
        wb = load_workbook(BytesIO(res.getvalue()))
        self.assertEqual(wb.sheetnames, SHEET_NAMES)

    def test_course_export_scoped(self):
        self.client.force_authenticate(self.outsider)
        res = self.client.get(reverse('analytics-course-export', kwargs={'pk': self.f.course.id}))
        self.assertEqual(res.status_code, status.HTTP_404_NOT_FOUND)

    def test_all_courses_export_uses_same_sheets(self):
        self.client.force_authenticate(self.f.creator)
        res = self.client.get(reverse('analytics-export'), {'ids': str(self.f.course.id)})
        wb = load_workbook(BytesIO(res.getvalue()))
        self.assertEqual(wb.sheetnames, SHEET_NAMES)
        self.assertEqual(len(rows(wb['Learners'])), 6)

    def test_all_courses_export_with_nothing_in_scope(self):
        self.client.force_authenticate(self.outsider)
        wb = load_workbook(BytesIO(self.client.get(reverse('analytics-export')).getvalue()))
        self.assertEqual(wb.sheetnames, SHEET_NAMES)
        self.assertEqual(rows(wb['Learners']), [])
