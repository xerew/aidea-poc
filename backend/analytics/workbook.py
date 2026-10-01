"""Learning analytics workbook — see the spec's "Excel workbook" section.

Long-format sheets (one row per learner × item, per visit, per event) for
statistical analysis, plus a wide Overview sheet. Seconds are plain numbers,
Overview uses minutes; timestamps are UTC ISO-8601 text. The all-courses
export uses the same sheets, with course columns on every row."""
from datetime import UTC

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from .learning import (
    INACTIVE_DAYS,
    STUCK_VISITS,
    CourseData,
    display_name,
    learner_row,
    resource_label,
    time_summary,
)

SHEET_NAMES = [
    'README', 'Overview', 'Learners', 'Modules', 'Activities', 'Resources',
    'Visits', 'Quiz answers', 'Events',
]

BASE = ['Course ID', 'Course', 'Learner ID', 'Learner']

HEADERS = {
    'Learners': BASE + [
        'Email', 'Enrolled at', 'Last active', 'Completed at', 'Progress %', 'Active s',
        'On-screen s', 'Visits', 'Status', 'Current position', 'Subject', 'Teaching level',
        'Role at school', 'Country', 'Interface language', 'Competency score',
        'Study participant (consented)',
    ],
    'Modules': BASE + [
        'Module ID', 'Module order', 'Module', 'Active s', 'On-screen s', 'Visits',
        'Activities done', 'Activities total', '% complete', 'First opened', 'Last seen',
        'Completed at',
    ],
    'Activities': BASE + [
        'Module ID', 'Module order', 'Module', 'Activity ID', 'Activity order', 'Activity',
        'Active s', 'On-screen s', 'Visits', 'Resources done', 'Resources total', 'Done',
        'First opened', 'Last seen', 'Completed at',
    ],
    'Resources': BASE + [
        'Module ID', 'Module order', 'Activity ID', 'Activity order', 'Activity',
        'Resource ID', 'Resource order', 'Resource', 'Type', 'Required', 'Active s',
        'On-screen s', 'Visits', 'First opened', 'Completed at', 'Quiz score %',
        'Video % watched', 'Scroll %', 'PDF opened', 'PDF downloaded', 'Image opened',
        'Assignment status',
    ],
    'Visits': BASE + [
        'Module ID', 'Module', 'Activity ID', 'Activity', 'Resource ID', 'Resource', 'Type',
        'Visit ID', 'Page visit', 'Started at', 'Last seen at', 'Active s', 'On-screen s',
        'Completed during', 'Language', 'Device', 'Local hour', 'TZ offset (min)',
        'Video % this visit',
    ],
    'Quiz answers': BASE + [
        'Module ID', 'Activity ID', 'Activity', 'Resource ID', 'Quiz', 'Question #',
        'Question', 'Option picked', 'Correct option', 'Right', 'Seconds on question',
        'Answered at',
    ],
    'Events': BASE + [
        'Resource ID', 'Resource', 'Type', 'Event', 'Time', 'Position s', 'From s', 'To s',
        'Question #', 'Option picked', 'Seconds on question',
    ],
}

README = [
    'AIDEA learning analytics export',
    'All times are UTC in ISO 8601 (…Z). Durations are seconds, except the Overview sheet (minutes).',
    'Active time: the page was visible and the learner scrolled, typed, clicked, touched or moved the pointer within the last 60 seconds, or a video was playing.',
    'On-screen time: the page was visible. Active time is never more than on-screen time.',
    'Each second of an activity page is credited to one resource: a playing video, else the resource last interacted with, else the one filling most of the screen.',
    'Module and activity times are sums of their resources. Visits = separate openings of an activity page.',
    'Typical time on the website = median among learners who opened the item and have measured time.',
    f'Status: completed = course finished; inactive = no activity for {INACTIVE_DAYS}+ days; '
    f'stuck = {STUCK_VISITS}+ visits to one resource without finishing it, or a quiz score below '
    'the pass threshold; on_track = otherwise. Checked in that order.',
    'Not tracked: work done before time tracking began has completion dates and scores but 0 visits and no time.',
    'Video % watched: share of the video actually played, in the best single visit.',
    "TZ offset: minutes ahead of UTC on the learner's device (Athens in summer = 180).",
    'Study participant (consented): the learner joined the AIDEA research study and gave consent.',
    'Sheets: Overview (one row per learner, wide), Learners, Modules, Activities, Resources '
    '(learner × item), Visits (raw), Quiz answers, Events (raw).',
]


def iso(dt):
    return dt.astimezone(UTC).strftime('%Y-%m-%dT%H:%M:%SZ') if dt else ''


def yes_no(value):
    return 'yes' if value else 'no'


def blank(value):
    return '' if value is None else value


def _minutes(seconds):
    return round(seconds / 60, 1)


def _base(cd, user):
    return [cd.course.id, cd.course.title, user.id, display_name(user)]


def _profile_columns(user):
    profile = getattr(user, 'profile', None)
    study = getattr(user, 'study', None)
    if profile is None:
        values = ['', '', '', '', '', '']
    else:
        values = [
            profile.subject.name if profile.subject else '',
            profile.teaching_level, profile.school_role, profile.country, profile.language,
            profile.competency_score,
        ]
    consented = bool(study and study.in_study and study.consented_at)
    return values + [yes_no(consented)]


def _learners(cd):
    for e in cd.enrollments:
        row = learner_row(cd, e)
        pos = row['position']
        position = f"{pos['module']} › {pos['activity']} › {pos['resource']}" if pos else ''
        yield _base(cd, e.user) + [
            e.user.email, iso(e.enrolled_at), iso(row['last_active']), iso(e.completed_at),
            e.progress_pct, row['active_s'], row['visible_s'], row['visits'], row['status'],
            position,
        ] + _profile_columns(e.user)


def _latest_completion(cd, uid, resources):
    times = [cd.completed_at(uid, r.id) for r in resources]
    return max((t for t in times if t), default=None)


def _modules(cd):
    for e in cd.enrollments:
        uid = e.user_id
        for m in cd.modules:
            t = time_summary(cd.module_visits(uid, m.id))
            activities = cd.activities_of.get(m.id, [])
            pct = cd.module_pct(uid, m.id)
            resources = [r for a in activities for r in cd.resources_of.get(a.id, [])]
            yield _base(cd, e.user) + [
                m.id, m.order, m.title, t['active_s'], t['visible_s'], t['visits'],
                sum(cd.activity_done(uid, a.id) for a in activities), len(activities), blank(pct),
                iso(t['first_opened']), iso(t['last_seen']),
                iso(_latest_completion(cd, uid, resources)) if pct == 100 else '',
            ]


def _activities(cd):
    for e in cd.enrollments:
        uid = e.user_id
        for m in cd.modules:
            for a in cd.activities_of.get(m.id, []):
                t = time_summary(cd.activity_visits(uid, a.id))
                resources = cd.resources_of.get(a.id, [])
                done = cd.activity_done(uid, a.id)
                yield _base(cd, e.user) + [
                    m.id, m.order, m.title, a.id, a.order, a.title, t['active_s'], t['visible_s'],
                    t['visits'], sum(cd.completed_at(uid, r.id) is not None for r in resources),
                    len(resources), yes_no(done), iso(t['first_opened']), iso(t['last_seen']),
                    iso(_latest_completion(cd, uid, resources)) if done else '',
                ]


def _resources(cd):
    for e in cd.enrollments:
        for r in cd.resources:
            d = cd.resource_detail(e.user_id, r)
            a = r.activity
            score = d['quiz_score']
            yield _base(cd, e.user) + [
                a.module_id, a.module.order, a.id, a.order, a.title, r.id, r.order,
                resource_label(r), r.type, yes_no(r.is_required), d['active_s'], d['visible_s'],
                d['visits'], iso(d['first_opened']), iso(d['completed_at']),
                round(score * 100) if score is not None else '', blank(d['video_pct']),
                blank(d['scroll_pct']), blank(d['pdf_opened']), blank(d['pdf_downloaded']),
                blank(d['image_opened']), blank(d['assignment_status']),
            ]


def _visits(cd):
    for e in cd.enrollments:
        for v in cd.user_visits(e.user_id):
            r = cd.resource_by_id.get(v.resource_id)
            if r is None:
                continue
            yield _base(cd, e.user) + [
                r.activity.module_id, r.activity.module.title, r.activity_id, r.activity.title,
                r.id, resource_label(r), r.type, v.id, str(v.page_key), iso(v.started_at),
                iso(v.last_seen_at), v.active_seconds, v.visible_seconds,
                yes_no(v.completed_during), v.language, v.device, blank(v.local_hour),
                blank(v.tz_offset_minutes), blank(v.media_progress.get('covered_pct')),
            ]


def _quiz_answers(cd):
    for e in cd.enrollments:
        for r in cd.resources:
            if r.type != 'quiz':
                continue
            for answer in cd.quiz_answers(e.user_id, r):
                right = answer['is_correct']
                yield _base(cd, e.user) + [
                    r.activity.module_id, r.activity_id, r.activity.title, r.id,
                    resource_label(r), answer['index'] + 1, answer['question'],
                    blank(answer['selected_text']), blank(answer['correct_text']),
                    '' if right is None else yes_no(right), blank(answer['seconds']),
                    iso(answer['answered_at']),
                ]


def _events(cd):
    for e in cd.enrollments:
        for r in cd.resources:
            for ev in cd.events_for(e.user_id, r.id):
                data = ev.data or {}
                index = data.get('question_index')
                yield _base(cd, e.user) + [
                    r.id, resource_label(r), r.type, ev.event_type, iso(ev.occurred_at),
                    blank(data.get('position')), blank(data.get('from')), blank(data.get('to')),
                    index + 1 if isinstance(index, int) else '', blank(data.get('selected')),
                    blank(data.get('seconds_on_question')),
                ]


ROWS = {
    'Learners': _learners, 'Modules': _modules, 'Activities': _activities,
    'Resources': _resources, 'Visits': _visits, 'Quiz answers': _quiz_answers,
    'Events': _events,
}


def _overview(ws, data):
    """Wide sheet: one row per learner; per module active minutes and % complete.
    With one course the module titles are in the headers; with several, the
    columns are by module position (titles are on the Modules sheet)."""
    width = max((len(cd.modules) for cd in data), default=0)
    single = data[0] if len(data) == 1 else None
    header = list(BASE)
    for i in range(width):
        label = f'M{i + 1} {single.modules[i].title}' if single else f'M{i + 1}'
        header += [f'{label} — active min', f'{label} — % complete']
    header += ['Total active min', 'Total on-screen min', 'Progress %', 'Status']
    _write_header(ws, header)
    for cd in data:
        for e in cd.enrollments:
            row = _base(cd, e.user)
            for i in range(width):
                if i < len(cd.modules):
                    m = cd.modules[i]
                    t = time_summary(cd.module_visits(e.user_id, m.id))
                    row += [_minutes(t['active_s']), blank(cd.module_pct(e.user_id, m.id))]
                else:
                    row += ['', '']
            t = time_summary(cd.user_visits(e.user_id))
            row += [_minutes(t['active_s']), _minutes(t['visible_s']), e.progress_pct, cd.status(e)]
            ws.append(row)


def _write_header(ws, header):
    ws.append(header)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    ws.freeze_panes = 'A2'
    for col in range(1, len(header) + 1):
        ws.column_dimensions[get_column_letter(col)].width = 16


def build_learning_workbook(courses, now=None):
    data = [CourseData(course, now) for course in courses]
    wb = Workbook()
    readme = wb.active
    readme.title = 'README'
    for line in README:
        readme.append([line])
    readme['A1'].font = Font(bold=True)
    readme.column_dimensions['A'].width = 120
    _overview(wb.create_sheet('Overview'), data)
    for name, rows in ROWS.items():
        ws = wb.create_sheet(name)
        _write_header(ws, HEADERS[name])
        for cd in data:
            for row in rows(cd):
                ws.append(row)
    return wb
